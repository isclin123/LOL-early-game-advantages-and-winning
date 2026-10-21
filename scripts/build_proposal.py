from pathlib import Path
import json
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'output'
OUT.mkdir(exist_ok=True)
(ROOT / 'audit').mkdir(exist_ok=True)

content = json.loads((ROOT / 'proposal_content.json').read_text(encoding='utf-8'))
title = content['title']
sections = content['sections']
questions = content['questions']
refs = content['refs']
author = content.get('author', '')
question_heading = content.get('question_heading', 'Ten Research Questions')
questions_after = content.get('questions_after', 'Data and Variables')
doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = sec.bottom_margin = Inches(0.72)
sec.left_margin = sec.right_margin = Inches(0.8)
for name in ['Normal', 'Title', 'Heading 1']:
    st = doc.styles[name]
    st.font.name = 'Times New Roman'
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.font.size = Pt(11)
normal = doc.styles['Normal'].paragraph_format
normal.line_spacing = 1.0
normal.space_after = Pt(5)
doc.styles['Title'].font.size = Pt(15)
doc.styles['Title'].font.bold = True
doc.styles['Title'].paragraph_format.space_after = Pt(5)
doc.styles['Title'].paragraph_format.line_spacing = 1.0
p = doc.add_paragraph(title, 'Title')
p.paragraph_format.keep_with_next = True
p = doc.add_paragraph('Statistics Capstone Project Proposal' + (' | ' + author if author else ''))
p.paragraph_format.space_after = Pt(10)
for r in p.runs:
    r.font.size = Pt(10)
for heading, body in sections:
    p = doc.add_paragraph()
    p.paragraph_format.widow_control = True
    p.add_run(heading + '. ').bold = True
    p.add_run(body)
    if heading == questions_after:
        p = doc.add_paragraph()
        p.add_run(question_heading).bold = True
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        for i, question in enumerate(questions, 1):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0 if i < len(questions) else 5)
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.first_line_indent = Inches(-0.25)
            p.add_run(f'{i}. ').bold = True
            p.add_run(question)
for ref in refs:
    p = doc.add_paragraph()
    source_text, source_url = ref.rsplit(' ', 1)
    p.add_run(source_text + ' ')
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), p.part.relate_to(source_url, RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    font = OxmlElement('w:rFonts')
    font.set(qn('w:ascii'), 'Times New Roman')
    font.set(qn('w:hAnsi'), 'Times New Roman')
    props.append(font)
    size = OxmlElement('w:sz')
    size.set(qn('w:val'), '18')
    props.append(size)
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '000000')
    props.append(color)
    run.append(props)
    text = OxmlElement('w:t')
    text.text = source_url
    run.append(text)
    link.append(run)
    p._p.append(link)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.0
    for r in p.runs:
        r.font.size = Pt(9)
# Remove any inherited borders from title styles.
for style in [doc.styles['Title'], doc.styles['Normal']]:
    ppr = style.element.find(qn('w:pPr'))
    if ppr is not None:
        borders = ppr.find(qn('w:pBdr'))
        if borders is not None:
            ppr.remove(borders)
doc.core_properties.title = title
doc.core_properties.subject = 'One-page semester statistics project proposal'
doc.core_properties.author = author or 'Sicheng Yi'
doc.save(OUT / 'LoL_Project_Proposal.docx')
md = '# ' + title + '\n\n'
if author:
    md += author + '\n\n'
for h, b in sections:
    md += '**' + h + '.** ' + b + '\n\n'
    if h == questions_after:
        md += '**' + question_heading + '**\n\n' + '\n'.join(f'{i}. {q}' for i, q in enumerate(questions, 1)) + '\n\n'
md += '\n\n' + '\n\n'.join(refs) + '\n'
(OUT / 'LoL_Project_Proposal.md').write_text(md, encoding='utf-8')
counts = {'body_words': len(re.findall(r"\b[\w]+(?:[-'][\w]+)*\b", ' '.join([b for _, b in sections] + (questions if questions_after else [])))),
          'title': title, 'sections': len(sections), 'research_questions': len(questions)}
(ROOT / 'audit' / 'proposal_content_check.json').write_text(json.dumps(counts, indent=2), encoding='utf-8')
print(json.dumps(counts))
